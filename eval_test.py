#!/usr/bin/env python3
"""RentSheet verification: structure + real formula evaluation (formulas engine)."""
import sys
from openpyxl import load_workbook

ok = 0; bad = 0
def check(name, cond):
    global ok, bad
    if cond: ok += 1; print(f"  ✓ {name}")
    else: bad += 1; print(f"  ✗ FAIL: {name}")

wb = load_workbook("dist/RentSheet.xlsx")
check("sheet order DASHBOARD first", wb.sheetnames[0] == "DASHBOARD")
check("sheets present", set(["DASHBOARD","README","PROPERTIES","RENT LOG","EXPENSES"]).issubset(set(wb.sheetnames)))
check("lease status formula", "EXPIRED" in (wb["PROPERTIES"]["F2"].value or ""))
check("rent balance formula", wb["RENT LOG"]["F2"].value == '=IF(B2="","",C2-D2)')
check("status formula has LATE", "LATE" in (wb["RENT LOG"]["G2"].value or ""))

import formulas
xl = formulas.ExcelModel().loads("dist/RentSheet.xlsx").finish()
sol = xl.calculate()
def get(sheet, cell):
    v = sol[f"'[RentSheet.xlsx]{sheet}'!{cell}"].value
    try: return float(v[0][0])
    except Exception:
        try: return float(v)
        except Exception: return v

# Sample data: rent log rows → 2A: 12000 paid + 6000 partial; 5B: 9500 paid
# collected = 27500; expected = 12000+12000+9500 = 33500; unpaid = 6000
check("collected = 27,500", get("DASHBOARD", "B6") == 27500)
check("expected = 33,500", get("DASHBOARD", "E6") == 33500)
check("unpaid balance = 6,000", get("DASHBOARD", "B9") == 6000)
check("collection rate ≈ 0.8209", abs(get("DASHBOARD", "E9") - 27500/33500) < 0.0001)
check("expenses = 850", get("DASHBOARD", "B12") == 850)
check("net income = 26,650", get("DASHBOARD", "E12") == 26650)
check("late/partial count = 1", get("DASHBOARD", "B15") == 1)
check("expired leases = 1 (5B lease ended 3 days ago)", get("DASHBOARD", "E15") == 1)
check("2A lease ACTIVE", "ACTIVE" in str(get("PROPERTIES", "F2")))
check("5B lease EXPIRED", "EXPIRED" in str(get("PROPERTIES", "F3")))
check("per-unit 2A collected = 18,000", abs(get("DASHBOARD", "C19") - 18000) < 0.001)

d = load_workbook("dist/RentSheet-Demo.xlsx")
check("demo watermark", any("DEMO VERSION" in str(c.value) for row in d["README"].iter_rows() for c in row if c.value))

print(f"\nRESULT: {ok} passed, {bad} failed")
sys.exit(1 if bad else 0)
